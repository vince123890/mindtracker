"""Generate Supabase seed SQL and scoring test fixture from the official tracker workbook.

Usage (from apps/):
    python scripts/generate_seed.py

Reads  ../docs/Dummy Tracker v1.4.xlsx
Writes supabase/seed.sql
       src/lib/scoring/fixtures/tracker-v1.4.json

Normalisation applied while reading (see docs/SPEK-RUMUS-Tracker-v1.4.md section 11):
  F-06  requirement blank but modifier > 0  -> modifier forced to 0
  F-07  deliverable code trimmed of spaces, NBSP and trailing dots
"""
import json
import random
import re
import sys
import uuid
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT.parent / "docs" / "Dummy Tracker v1.4.xlsx"
SEED = ROOT / "supabase" / "seed.sql"
FIXTURE = ROOT / "src" / "lib" / "scoring" / "fixtures" / "tracker-v1.4.json"

PHASES = ["FEL-0", "FEL-1", "FEL-2", "FEL-3"]
TYPES = ["NFP", "NMF", "MPM", "m-PM", "NSI"]  # order of columns J:N and O:S

CHAPTERS = {  # chapter_no -> (code prefix, name, dimension_id)
    1: ("E", "Exploration", None),
    2: ("MP", "Mining", None),
    3: ("GM", "Geo-Metallurgy & Process", 2),
    4: ("E", "Engineering", 2),
    5: ("PC", "Procurement & Construction", None),
    6: ("TM", "Tailings Management", 2),
    7: ("WM", "Water Management", 2),
    8: ("INF", "Infrastructure Activities", None),
    9: ("GPL", "Government Relations & Permits", None),
    10: ("BC", "Business & Commercial", 1),
    11: ("RM", "Risk Management", 4),
    12: ("ES", "Environmental & Social", None),
    13: ("OR", "Operational Readiness", None),
    14: ("OHSS/PSM", "OHSS & Process Safety Management", None),
    15: ("PM", "Project Governance & Monitoring", 3),
}
DIMENSIONS = {1: "Business and Commercials", 2: "Resource and Technical", 3: "Project Execution", 4: "Risk Management"}


def q(v):
    if v is None:
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return repr(round(v, 6)) if isinstance(v, float) else str(v)
    return "'" + str(v).replace("'", "''") + "'"


def clean_code(raw):
    return re.sub(r"[\s\xa0.]+$", "", str(raw)).strip()


def read_tracker():
    wb = openpyxl.load_workbook(XLSX)
    cached = openpyxl.load_workbook(XLSX, data_only=True)
    deliverables, warnings = [], []
    excel_cache = {}
    next_id = 1
    for phase in PHASES:
        ws, wc = wb[phase], cached[phase]
        for r in range(10, ws.max_row + 1):
            code = ws.cell(r, 3).value
            if not code:
                continue
            chapter_no = int(str(ws.cell(r, 2).value).split()[-1])
            appl = {}
            for i, t in enumerate(TYPES):
                marker = ws.cell(r, 10 + i).value  # J:N
                modifier = float(ws.cell(r, 15 + i).value or 0)  # O:S
                marker = {"G": "G", "A": "A", "X": "A", "C": "C"}.get(marker, "N")
                if marker == "N" and modifier > 0:
                    warnings.append(f"F-06 {phase} row {r} {t}: requirement blank, modifier {modifier} -> 0")
                    modifier = 0.0
                if marker != "N" and modifier == 0:
                    warnings.append(f"F-06 {phase} row {r} {t}: requirement {marker} with modifier 0 (kept)")
                appl[t] = {"marker": marker, "modifier": modifier}
            d = {
                "id": next_id,
                "phase": phase,
                "excelRow": r,
                "chapterNo": chapter_no,
                "code": clean_code(code),
                "name": str(ws.cell(r, 4).value),
                "referenceDocument": ws.cell(r, 5).value,
                "baseWeight": float(ws.cell(r, 9).value),
                "applicability": appl,
                "excelScore": ws.cell(r, 25).value,
            }
            deliverables.append(d)
            next_id += 1
        excel_cache[phase] = {
            "index1": wc.cell(3, 5).value,
            "index2": wc.cell(4, 5).value,
            "index3": wc.cell(5, 5).value,
            "index4": wc.cell(6, 5).value,
        }
    return deliverables, excel_cache, warnings


# ---------------------------------------------------------------- scoring (mirror of src/lib/scoring, production mode)
STRENGTH = {"G": 3, "A": 2, "C": 1, "N": 0}


def accepted(d, t1, t2, resolution=None):
    m1 = d["applicability"][t1]["marker"]
    m2 = d["applicability"][t2]["marker"] if t2 else "N"
    overall = m1 if STRENGTH[m1] >= STRENGTH[m2] else m2
    if overall == "C":
        return {"GATE": "G", "APPLICABLE": "A", "NOT_REQUIRED": "N"}.get(resolution, "?")
    return overall


def eff_weight(d, t1, t2):
    m1 = d["applicability"][t1]["modifier"]
    m2 = d["applicability"][t2]["modifier"] if t2 else 0
    return d["baseWeight"] * max(m1, m2)


def indices(rows):
    num = den = 0.0
    g_total = g_done = started = counted = pending = 0
    for acc, w, score in rows:
        if acc == "?":
            pending += 1
            continue
        if acc not in ("G", "A"):
            continue
        counted += 1
        den += 4 * w
        num += w * (score or 0)
        if (score or 0) >= 1:
            started += 1
        if acc == "G":
            g_total += 1
            if (score or 0) >= 3:
                g_done += 1
    return {
        "index1": round(num / den, 4) if den else None,
        "index2": round(g_done / g_total, 4) if g_total else None,
        "index3": round(started / counted, 4) if counted else None,
        "index4": pending,
    }


# ---------------------------------------------------------------- demo data
ORGS = [
    ("MIND", "PT Mineral Industri Indonesia (Persero)", "HOLDING"),
    ("ANTAM", "PT Aneka Tambang Tbk", "MEMBER"),
    ("PTBA", "PT Bukit Asam Tbk", "MEMBER"),
    ("TIMAH", "PT Timah Tbk", "MEMBER"),
    ("INALUM", "PT Indonesia Asahan Aluminium", "MEMBER"),
]

PROJECTS = [
    # id, code, name, org, location, t1, t2, entry, phase plan {phase: (status, profile)}
    ("11111111-1111-4111-8111-000000000001", "PRJ-ANTAM-001", "Smelter Feronikel Halmahera Timur", "ANTAM",
     "Halmahera Timur, Maluku Utara", "NMF", "NFP", "FEL-0",
     [("FEL-0", "APPROVED", "pass"), ("FEL-1", "DRAFT", "partial")]),
    ("11111111-1111-4111-8111-000000000002", "PRJ-ANTAM-002", "Pabrik Alumina Tahap 2", "ANTAM",
     "Mempawah, Kalimantan Barat", "MPM", None, "FEL-1",
     [("FEL-1", "WAITING_APPROVAL", "pass")]),
    ("11111111-1111-4111-8111-000000000003", "PRJ-TIMAH-001", "Smelter Timah Muntok", "TIMAH",
     "Muntok, Bangka Barat", "MPM", "m-PM", "FEL-1",
     [("FEL-1", "DRAFT", "early")]),
    ("11111111-1111-4111-8111-000000000004", "PRJ-PTBA-001", "Hilirisasi Batubara (DME)", "PTBA",
     "Tanjung Enim, Sumatera Selatan", "NSI", "NMF", "FEL-2",
     [("FEL-2", "DRAFT", "partial")]),
]

USERS = {
    "TIM": {"ANTAM": "u-tim-antam"},
    "PMO_AH": {"ANTAM": "u-pmo-antam", "TIMAH": "u-pmo-timah"},
}


def demo_scores(rng, profile, acc):
    """Return (score, is_na) for a deliverable with accepted requirement acc."""
    if acc not in ("G", "A"):
        return None, False
    if profile == "pass":
        return (rng.choice([3, 3, 4, 4]) if acc == "G" else rng.choice([2, 3, 3, 4, 4])), False
    if profile == "partial":
        roll = rng.random()
        if roll < 0.15:
            return None, False
        if acc == "G":
            return rng.choice([2, 3, 3, 4]), False
        return rng.choice([1, 2, 3, 3, 4]), False
    # early
    roll = rng.random()
    if roll < 0.55:
        return None, False
    return rng.choice([0, 1, 1, 2, 2, 3]), False


def build_seed(deliverables):
    rng = random.Random(14)
    out = []
    w = out.append
    w("-- Generated by scripts/generate_seed.py from docs/Dummy Tracker v1.4.xlsx. Do not edit by hand.")
    w("begin;")
    w("insert into organization (id, name, kind) values " + ",\n  ".join(
        f"({q(a)}, {q(b)}, {q(c)})" for a, b, c in ORGS) + ";")
    users = [
        ("u-pmo-admin", "Rina (PMO MIND ID)", "PMO_ADMIN", "MIND"),
        ("u-dir-mind", "Budi (Direktur MIND ID)", "DIREKTUR_MIND_ID", "MIND"),
        ("u-div-mind", "Sari (Divisi MIND ID)", "DIVISI_MIND_ID", "MIND"),
        ("u-div-do", "Andi (Divisi DO MIND ID)", "DIVISI_DO_MIND_ID", "MIND"),
        ("u-dir-antam", "Dewi (Direktur Antam)", "DIREKTUR_AH", "ANTAM"),
        ("u-pmo-antam", "Eko (PMO Antam)", "PMO_AH", "ANTAM"),
        ("u-tim-antam", "Fajar (Tim Proyek Antam)", "TIM_PROYEK", "ANTAM"),
        ("u-pic-antam", "Gita (PIC Operasi Antam)", "PIC_OPERASI_AH", "ANTAM"),
        ("u-pmo-timah", "Hadi (PMO Timah)", "PMO_AH", "TIMAH"),
    ]
    w("insert into app_user (id, name, role, organization_id) values " + ",\n  ".join(
        f"({q(a)}, {q(b)}, {q(c)}, {q(d)})" for a, b, c, d in users) + ";")

    w("insert into project_type (code, name, modifier_header, sort_order) values "
      "('NFP','New Fixed Plant Facilities','xNFP',1),('NMF','New Mining Facilities','xNMF',2),"
      "('MPM','Major Plant Modifications','xMPM',3),('m-PM','minor Plant Modifications','xm-PM',4),"
      "('NSI','New Supporting Infrastructures','xNSI',5);")
    w("insert into phase (code, name, seq, is_active) values ('FEL-0','FEL-0',1,true),('FEL-1','FEL-1',2,true),"
      "('FEL-2','FEL-2',3,true),('FEL-3','FEL-3',4,true),('DESIGN','Design',5,false),"
      "('CONSTRUCTION','Construction',6,false),('COMMISSIONING','Commissioning',7,false);")
    w("insert into dimension (id, name, sort_order) values " + ",".join(
        f"({k},{q(v)},{k})" for k, v in DIMENSIONS.items()) + ";")
    w("insert into chapter_function (chapter_no, code_prefix, name, dimension_id) values " + ",\n  ".join(
        f"({n},{q(p)},{q(nm)},{q(dim)})" for n, (p, nm, dim) in CHAPTERS.items()) + ";")
    w("insert into scoring_config (version, start_threshold, pass_threshold, max_score, gate_min_maturity_pct, "
      "gate_control_must_complete, na_treatment, created_by) values (1,1,3,4,0.60,true,'EXCLUDE','u-pmo-admin');")

    rows = []
    appl = []
    for d in deliverables:
        rows.append(f"({d['id']},{q(d['phase'])},{d['chapterNo']},{q(d['code'])},{q(d['name'])},"
                    f"{q(d['referenceDocument'])},{d['baseWeight']},{d['excelRow']})")
        for t in TYPES:
            a = d["applicability"][t]
            appl.append(f"({d['id']},{q(t)},{q(a['marker'])},{a['modifier']})")
    for i in range(0, len(rows), 200):
        w("insert into deliverable (id, phase_code, chapter_no, code, name, reference_document, base_weight, sort_order) values\n  "
          + ",\n  ".join(rows[i:i + 200]) + ";")
    for i in range(0, len(appl), 500):
        w("insert into deliverable_applicability (deliverable_id, project_type_code, marker, modifier) values\n  "
          + ",\n  ".join(appl[i:i + 500]) + ";")
    w(f"select setval(pg_get_serial_sequence('deliverable','id'), {len(deliverables)});")

    by_phase = {}
    for d in deliverables:
        by_phase.setdefault(d["phase"], []).append(d)

    for pid, code, name, org, loc, t1, t2, entry, plan in PROJECTS:
        current = plan[-1][0]
        w("insert into project (id, code, name, organization_id, location, project_type_1, project_type_2, "
          "entry_phase, current_phase, project_manager, executive_sponsor, description, created_by) values "
          f"({q(pid)},{q(code)},{q(name)},{q(org)},{q(loc)},{q(t1)},{q(t2)},{q(entry)},{q(current)},"
          f"'Ir. Project Manager {org}','Direktur Pengembangan {org}','Proyek demo prototype','u-pmo-admin');")
        for phase, status, profile in plan:
            ppi = str(uuid.uuid5(uuid.NAMESPACE_URL, f"{pid}/{phase}"))
            submitted_by = "u-pmo-antam" if org == "ANTAM" else "u-pmo-admin"
            submitted = status in ("WAITING_APPROVAL", "APPROVED")
            decided = status == "APPROVED"
            comment = {"pass": "Seluruh deliverable Gate Control lengkap; siap ditinjau.",
                       "partial": "Beberapa deliverable Gate Control masih di bawah skor 3.",
                       "early": "Fase baru dimulai, pengisian berjalan."}[profile]
            w("insert into project_phase_instance (id, project_id, phase_code, gate_status, submitted_by, submitted_at, "
              "decided_by, decided_at, summary_comment) values "
              f"({q(ppi)},{q(pid)},{q(phase)},{q(status)},{q(submitted_by if submitted else None)},"
              f"{'now() - interval ' + q('3 days') if submitted else 'null'},"
              f"{q('u-pmo-admin' if decided else None)},{'now() - interval ' + q('1 day') if decided else 'null'},"
              f"{q(comment)});")
            score_rows, calc = [], []
            for d in by_phase[phase]:
                ov_acc = accepted(d, t1, t2)
                resolution = justification = None
                if ov_acc == "?" and (profile != "early" or rng.random() < 0.5):
                    resolution = rng.choice(["APPLICABLE", "APPLICABLE", "GATE", "NOT_REQUIRED"])
                    if profile == "partial" and rng.random() < 0.08:
                        resolution = None
                    if resolution:
                        justification = "Diputuskan bersama tim proyek pada rapat koordinasi fase."
                acc = accepted(d, t1, t2, resolution)
                score, is_na = demo_scores(rng, profile, acc)
                calc.append((acc, eff_weight(d, t1, t2), score))
                user = USERS["TIM"].get(org, "u-pmo-admin")
                score_rows.append(
                    f"({q(ppi)},{d['id']},{q(score)},{q(is_na)},{q(resolution)},{q(justification)},"
                    f"{q(USERS['PMO_AH'].get(org, 'u-pmo-admin') if resolution else None)},"
                    f"{q('Bukti terlampir pada folder proyek.' if score is not None and score >= 3 else None)},"
                    f"{q(user if score is not None else None)})")
            w("insert into deliverable_score (phase_instance_id, deliverable_id, score, is_na, conditional_resolution, "
              "conditional_justification, conditional_decided_by, notes, updated_by) values\n  "
              + ",\n  ".join(score_rows) + ";")
            idx = indices(calc)
            if status == "DRAFT":
                # Riwayat snapshot bulanan (Progress Curve): indeks dihitung ulang dari sebagian skor
                # yang sudah terisi pada akhir bulan tersebut — pola "menumpuk di akhir" sengaja diperagakan.
                for label, frac in (("2026-07", 0.25), ("2026-08", 0.4), ("2026-09", 0.85)):
                    hist = [(a, wt, sc if (sc is not None and rng.random() < frac) else None) for a, wt, sc in calc]
                    w("insert into score_snapshot (phase_instance_id, label, indices, config_version) values "
                      f"({q(ppi)},{q(label)},{q(json.dumps(indices(hist)))},1);")
            if submitted:
                w("insert into phase_gate_transition (phase_instance_id, from_status, to_status, actor_id, note, indices) values "
                  f"({q(ppi)},'DRAFT','WAITING_APPROVAL',{q(submitted_by)},null,{q(json.dumps(idx))});")
            if decided:
                w("insert into phase_gate_transition (phase_instance_id, from_status, to_status, actor_id, note, indices) values "
                  f"({q(ppi)},'WAITING_APPROVAL','APPROVED','u-pmo-admin','Disetujui untuk naik fase.',{q(json.dumps(idx))});")
                w("insert into score_snapshot (phase_instance_id, label, indices, config_version) values "
                  f"({q(ppi)},'PHASE_CLOSE',{q(json.dumps(idx))},1);")
            print(f"  {code} {phase} {status}: {idx}", file=sys.stderr)

    # feedback (revision log) on the partial FEL-1 project
    ppi = str(uuid.uuid5(uuid.NAMESPACE_URL, f"{PROJECTS[0][0]}/FEL-1"))
    fb = next(d for d in by_phase["FEL-1"] if accepted(d, "NMF", "NFP") == "G")
    w("insert into revision_log_entry (phase_instance_id, deliverable_id, author_id, body, is_blocking) values "
      f"({q(ppi)},{fb['id']},'u-div-mind','Mohon lampirkan studi pasar terbaru sebagai dasar skor.',false);")

    # --- Dashboard DO dummy (pengganti MCT) ---
    plants = [("ANTAM", "Smelter Pomalaa", "Feronikel", "ton Ni"), ("ANTAM", "Refinery Logam Mulia", "Emas", "kg"),
              ("PTBA", "Tambang Tanjung Enim", "Batubara", "ton"), ("TIMAH", "Smelter Muntok", "Timah", "ton Sn"),
              ("INALUM", "Smelter Kuala Tanjung", "Aluminium", "ton Al")]
    prod = []
    for org, plant, product, unit in plants:
        base = {"Feronikel": 2100, "Emas": 95, "Batubara": 2_400_000, "Timah": 1500, "Aluminium": 21000}[product]
        for m in range(1, 7):
            target = base
            actual = round(base * rng.uniform(0.82, 1.06), 2)
            prod.append(f"({q(org)},{q(plant)},{q(product)},{q(unit)},'2026-{m:02d}-01',{target},{actual})")
    w("insert into do_production (organization_id, plant, product, unit, period, target, actual) values\n  "
      + ",\n  ".join(prod) + ";")
    w("insert into rca (id, number, organization_id, plant, product, unit, period, gap, category, problem, "
      "root_cause_category, root_cause, analysis, status, created_by) values "
      "('22222222-2222-4222-8222-000000000001','RCA-ANTAM-2026-001','ANTAM','Smelter Pomalaa','Feronikel','ton Ni',"
      "'2026-04-01',-310,'Teknis','Realisasi produksi April di bawah target RKAP.','Machine',"
      "'Gangguan electric furnace #2 selama 9 hari.','Unplanned shutdown akibat kerusakan elektroda.','DISETUJUI','u-pic-antam'),"
      "('22222222-2222-4222-8222-000000000002','RCA-ANTAM-2026-002','ANTAM','Smelter Pomalaa','Feronikel','ton Ni',"
      "'2026-05-01',-180,'Bahan Baku','Kadar bijih masuk lebih rendah dari rencana.','Material',"
      "'Pasokan bijih kadar 1,4% dari pit baru.','Blending belum optimal.','MENUNGGU_REVIEW','u-pic-antam');")
    w("insert into rca_review_log (rca_id, from_status, to_status, actor_id, note) values "
      "('22222222-2222-4222-8222-000000000001','DRAFT','MENUNGGU_REVIEW','u-pic-antam',null),"
      "('22222222-2222-4222-8222-000000000001','MENUNGGU_REVIEW','DISETUJUI','u-div-do','Analisis memadai.'),"
      "('22222222-2222-4222-8222-000000000002','DRAFT','MENUNGGU_REVIEW','u-pic-antam',null);")
    w("insert into action_plan (rca_id, action, pic, target_date, progress, status, created_by) values "
      "('22222222-2222-4222-8222-000000000001','Penggantian elektroda dan inspeksi rutin mingguan','Kepala Maintenance Pomalaa','2026-06-30',100,'VERIFIED','u-pic-antam'),"
      "('22222222-2222-4222-8222-000000000001','Penyusunan SOP predictive maintenance furnace','Superintendent Furnace','2026-08-31',60,'IN_PROGRESS','u-pic-antam');")

    # --- Strategy & Simulation dummy (pengganti aplikasi simulasi eksternal) ---
    w("insert into strategy_scenario (code, name, description, source, variables, projections) values "
      "('SCN-01','Baseline RKAP 2027','Asumsi harga & volume sesuai RKAP','Simulasi Eksternal (dummy)',"
      "'{\"harga_nikel_usd_t\":16500,\"kurs\":15800,\"utilisasi\":0.85}','{\"pendapatan_t_rp\":12.4,\"ebitda_t_rp\":3.1,\"irr\":0.14}'),"
      "('SCN-02','Harga Nikel Turun 15%','Sensitivitas penurunan harga','Simulasi Eksternal (dummy)',"
      "'{\"harga_nikel_usd_t\":14000,\"kurs\":15800,\"utilisasi\":0.85}','{\"pendapatan_t_rp\":10.6,\"ebitda_t_rp\":2.0,\"irr\":0.10}'),"
      "('SCN-03','Ekspansi Kapasitas','Penambahan line smelter','Simulasi Eksternal (dummy)',"
      "'{\"harga_nikel_usd_t\":16500,\"kurs\":15800,\"utilisasi\":0.92}','{\"pendapatan_t_rp\":15.2,\"ebitda_t_rp\":3.9,\"irr\":0.16}');")
    build_rto(w, rng)
    build_dummy_sources(w, rng)
    w("commit;")
    return "\n".join(out) + "\n"


RTO_PILLARS = [
    (1, "Leadership Readiness", [
        "Business Objectives and Operating Philosophy", "Health and Safety Commitment", "ESG Commitment",
        "Culture and Core Values", "Operational Readiness Leadership and Team", "Project Commitment"]),
    (2, "Organization & People Readiness", [
        "Business Process and RACI", "Organizational Structure", "Internal Policies and Document Control",
        "Staffing for Commissioning, Startup and Operations", "Competency and Development Program",
        "Communications and Workflow Management"]),
    (3, "Operations Readiness", None),
    (4, "Maintenance Readiness", None),
    (5, "Process Readiness", None),
    (6, "Support Readiness", None),
]


def build_rto(w, rng):
    """Master RTO: sub-elemen pilar 1-2 terbaca dari RTO Tool v0.6; pilar 3-6 & teks requirement = placeholder (OI-02)."""
    desc = "Deskripsi pilar dari worksheet Descriptions (menunggu berkas RTO penuh, OI-02)."
    w("insert into rto_pillar (id, name, description) values " + ",".join(
        f"({pid},{q(name)},{q(desc)})" for pid, name, _ in RTO_PILLARS) + ";")
    subs, reqs = [], []
    sub_id = req_id = 0
    pillar_of_req = {}
    for pid, _, names in RTO_PILLARS:
        verified = names is not None
        names = names or [f"Sub-elemen {pid}.{i} (placeholder, OI-02)" for i in range(1, 4)]
        for i, nm in enumerate(names, 1):
            sub_id += 1
            subs.append(f"({sub_id},{pid},{q(f'{pid}.{i}')},{q(nm)},{q(verified)})")
            for k in range(1, 4):
                req_id += 1
                code = f"{pid}.{i}.{k}"
                pillar_of_req[req_id] = pid
                reqs.append(f"({req_id},{sub_id},{q(code)},"
                            f"{q(f'Requirement {code} - teks deliverable menunggu berkas RTO (OI-02)')},"
                            f"{q('Panduan pemenuhan dari worksheet ProgressGuidance (placeholder).')},1)")
    w("insert into rto_sub_element (id, pillar_id, code, name, verified) values " + ",\n  ".join(subs) + ";")
    w("insert into rto_requirement (id, sub_element_id, code, deliverable_description, guidance, weight) values\n  "
      + ",\n  ".join(reqs) + ";")
    w(f"select setval(pg_get_serial_sequence('rto_sub_element','id'), {sub_id});")
    w(f"select setval(pg_get_serial_sequence('rto_requirement','id'), {req_id});")
    pid = PROJECTS[0][0]
    w(f"insert into rto_project_sub_element (project_id, sub_element_id, applicable) select {q(pid)}, id, "
      "case when code in ('1.3','5.3') then false else true end from rto_sub_element;")
    aid = "33333333-3333-4333-8333-000000000001"
    w("insert into rto_assessment (id, project_id, assessment_date, assessment_point, led_by, status, created_by) values "
      f"({q(aid)},{q(pid)},'2026-09-15','FEL-1 - Readiness review awal','OR Team Lead Antam','DRAFT','u-pmo-antam');")
    res = []
    for rid in range(1, req_id + 1):
        pillar = pillar_of_req[rid]
        if pillar <= 2:
            score = rng.choice([2, 3, 3, 4, 4])
            delivered = rng.random() < 0.7
        elif pillar <= 4:
            score = rng.choice([None, 1, 2, 2, 3])
            delivered = score is not None and rng.random() < 0.4
        else:
            score, delivered = None, False
        res.append(f"({q(aid)},{rid},{q(score)},{q(delivered)},'u-tim-antam')")
    w("insert into rto_result (assessment_id, requirement_id, maturity_score, is_delivered, updated_by) values\n  "
      + ",\n  ".join(res) + ";")


def build_dummy_sources(w, rng):
    """Seed pengganti integrasi I-1, I-2, I-3, I-4 - DIGANTI integrasi nyata saat development."""
    w("insert into ext_project_progress (project_code, stage, physical_progress, schedule_status, target_cod) values "
      "('PRJ-ANTAM-001','FEL-1',8.5,'ON_TRACK','2029-12-31'),('PRJ-ANTAM-002','FEL-1',5.0,'AT_RISK','2029-06-30'),"
      "('PRJ-TIMAH-001','FEL-1',3.0,'ON_TRACK','2030-03-31'),('PRJ-PTBA-001','FEL-2',22.0,'DELAYED','2028-12-31'),"
      "('PRJ-INALUM-EPC','EPC',61.0,'AT_RISK','2027-09-30');")
    w("insert into ext_project_risk (project_code, risk_rating, open_risks, mitigation_on_track, top_risk) values "
      "('PRJ-ANTAM-001','MEDIUM',14,78,'Ketersediaan pasokan listrik'),('PRJ-ANTAM-002','HIGH',21,55,'Pembebasan lahan'),"
      "('PRJ-TIMAH-001','LOW',6,90,'Kadar bijih'),('PRJ-PTBA-001','EXTREME',27,41,'Kelayakan offtaker DME'),"
      "('PRJ-INALUM-EPC','HIGH',18,63,'Keterlambatan pengadaan peralatan');")
    w("insert into ext_project_document (project_code, rkap_status, fs_completeness, fid_status) values "
      "('PRJ-ANTAM-001','Masuk RKAP 2027',45,'Belum'),('PRJ-ANTAM-002','Masuk RKAP 2027',60,'Belum'),"
      "('PRJ-TIMAH-001','Usulan RKAP 2027',20,'Belum'),('PRJ-PTBA-001','Masuk RKAP 2026',85,'Dalam evaluasi'),"
      "('PRJ-INALUM-EPC','Masuk RKAP 2025',100,'Disetujui');")
    params = [("ANTAM", "Smelter Pomalaa", "Recovery Ni", "%", 92.0, True),
              ("ANTAM", "Smelter Pomalaa", "Konsumsi listrik", "MWh/t Ni", 38.0, False),
              ("TIMAH", "Smelter Muntok", "Recovery Sn", "%", 96.5, True),
              ("PTBA", "Tambang Tanjung Enim", "Stripping ratio", "bcm/t", 4.5, False),
              ("INALUM", "Smelter Kuala Tanjung", "Current efficiency", "%", 93.0, True)]
    rows = []
    for org, plant, par, unit, target, hib in params:
        for m in range(1, 7):
            actual = round(target * rng.uniform(0.93, 1.05), 2)
            rows.append(f"({q(org)},{q(plant)},{q(par)},{q(unit)},'2026-{m:02d}-01',{target},{actual},{q(hib)})")
    w("insert into do_key_parameter (organization_id, plant, parameter, unit, period, target, actual, higher_is_better) values\n  "
      + ",\n  ".join(rows) + ";")
    eq = [("ANTAM", "Smelter Pomalaa", "Electric Furnace #2"), ("ANTAM", "Smelter Pomalaa", "Rotary Kiln #1"),
          ("TIMAH", "Smelter Muntok", "Reverberatory Furnace"), ("PTBA", "Tambang Tanjung Enim", "Excavator PC2000"),
          ("INALUM", "Smelter Kuala Tanjung", "Pot Line 3")]
    rows = []
    for org, plant, e in eq:
        for m in range(1, 7):
            rows.append(f"({q(org)},{q(plant)},{q(e)},'2026-{m:02d}-01',{round(rng.uniform(84, 98), 1)},"
                        f"{round(rng.uniform(3, 18), 1)},{round(rng.uniform(150, 720), 0)},{rng.randint(2, 25)})")
    w("insert into do_maintenance (organization_id, plant, equipment, period, availability, mttr_hours, mtbf_hours, open_work_orders) values\n  "
      + ",\n  ".join(rows) + ";")
    w("insert into gate_workflow_step (step_order, name, actor_role, sla_days, description) values "
      "(1,'Pengisian & pengajuan','PMO_AH',0,'Tim proyek mengisi skor; PMO AH mengajukan gate bila Index 4 = 0, seluruh G >= 3, Index 1 >= 60%'),"
      "(2,'Review & keputusan','PMO_ADMIN',5,'PMO MIND ID menyetujui atau meminta revisi; pengaju tidak boleh menyetujui');")


def main():
    deliverables, excel_cache, warnings = read_tracker()
    for line in warnings:
        print(line, file=sys.stderr)
    SEED.write_text(build_seed(deliverables), encoding="utf-8")
    fixture = {
        "source": "docs/Dummy Tracker v1.4.xlsx",
        "projectTypes": {"type1": "NMF", "type2": "NFP"},
        "excelCachedIndices": excel_cache,
        "deliverables": [
            {k: d[k] for k in ("id", "phase", "excelRow", "chapterNo", "code", "baseWeight", "applicability", "excelScore")}
            for d in deliverables
        ],
    }
    FIXTURE.write_text(json.dumps(fixture, indent=1), encoding="utf-8")
    print(f"deliverables={len(deliverables)} seed={SEED} fixture={FIXTURE}", file=sys.stderr)


if __name__ == "__main__":
    main()
